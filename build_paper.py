"""Build the editable Word preprint with native mathematical equations."""
from pathlib import Path
import json,re,argparse
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

def el(tag,**attrs):
    x=OxmlElement(tag)
    for k,v in attrs.items():x.set(qn(k),str(v))
    return x
def mr(text):
    x=el('m:r');t=el('m:t');t.text=text;t.set(qn('xml:space'),'preserve');x.append(t);return x
def seq(*parts):
    items=[]
    for p in parts:items.extend(p if isinstance(p,list) else [mr(p) if isinstance(p,str) else p])
    return items
def box(tag,parts):
    x=el(tag)
    for p in seq(parts):x.append(p)
    return x
def sub(base,index):
    x=el('m:sSub');x.append(box('m:e',base));x.append(box('m:sub',index));return x
def sup(base,index):
    x=el('m:sSup');x.append(box('m:e',base));x.append(box('m:sup',index));return x
def frac(a,b):
    x=el('m:f');x.append(box('m:num',a));x.append(box('m:den',b));return x
def summation(lower,expr):
    x=el('m:nary');pr=el('m:naryPr');pr.append(el('m:chr',**{'m:val':'∑'}));pr.append(el('m:limLoc',**{'m:val':'subSup'}));pr.append(el('m:supHide',**{'m:val':'1'}));x.append(pr);x.append(box('m:sub',lower));x.append(el('m:sup'));x.append(box('m:e',expr));return x
def eq(number):
    if number=='1':return seq(sub('x','i'),'(t) = θ(t) + ',sub('b','i'),' + ',sub('ε','i'),'(t),     ',sub('L','s'),'(t) = ',sup(seq('‖',sub('p','s'),'(t) − θ(t)‖'),'2'))
    if number=='2':return seq(sub('R','i'),'(t+1) = α',sub('R','i'),'(t) − (1−α)',summation('k=1,…,q',sup(seq('[',sub('x','ik'),'(t) − ',sub('y','k'),'(t)]'),'2')),'   (δ=0)')
    if number=='3':return seq(sub('w','i'),' = ',frac(seq('exp(',sub('R','i'),')'),summation('j∈C',seq('exp(',sub('R','j'),')'))),',     ',sub('N','i'),' = {j∈C : ‖',sub('x','j'),'−',sub('x','i'),'‖ ≤ ',sub('τ','i'),'}')
    if number=='4':return seq(sub('z','i'),' = (1−ρ)',sub('x','i'),' + ρ',frac(summation('j∈Nᵢ',seq('exp(',sub('R','j'),')',sub('x','j'))),summation('j∈Nᵢ',seq('exp(',sub('R','j'),')'))),',     p = ',summation('i∈C',seq(sub('w','i'),sub('z','i'))))
    if number=='5':return seq(summation('i∈C',seq(sub('w','i'),'[(1−ρ)',sub('x','i'),' + ρp]')),' = (1−ρ)p + ρp = p')
    if number=='6':return seq('E',sup('‖x̄−θ‖','2'),' = ',sup(seq('‖',frac('1','N'),summation('i',sub('b','i')),'‖'),'2'),' + ',frac('1',sup('N','2')),summation('i',seq('tr(',sub('Σ','i'),')')))
    raise ValueError(number)
def hyperlink(p,url,label=None):
    rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True);h=el('w:hyperlink',**{'r:id':rel});r=el('w:r');pr=el('w:rPr');pr.append(el('w:color',**{'w:val':'333333'}));r.append(pr);t=el('w:t');t.text=label or url;r.append(t);h.append(r);p._p.append(h)
def inline(p,text):
    for part in re.split(r'(https?://\S+)',text):
        if part.startswith('http'):hyperlink(p,part.rstrip('.')); p.add_run('.' if part.endswith('.') else '')
        else:
            for token in re.split(r'([xbyR]_i|σ_y)',part):
                if re.fullmatch(r'[xbyR]_i|σ_y',token):
                    base,index=token.split('_');p.add_run(base);p.add_run(index).font.subscript=True
                else:p.add_run(token)

doc=Document();section=doc.sections[0]
section.page_width=Inches(8.5);section.page_height=Inches(11)
section.top_margin=Inches(.72);section.bottom_margin=Inches(.72)
section.left_margin=Inches(.85);section.right_margin=Inches(.85)
section.footer_distance=Inches(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='Cambria';st.font.color.rgb=RGBColor(0,0,0)
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.08;normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.widow_control=True
doc.styles['Title'].font.size=Pt(21);doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(12)
for name,size in [('Heading 1',14),('Heading 2',12)]:
    st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(12);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(9.5);doc.styles['Caption'].font.italic=False;doc.styles['Caption'].font.bold=False
doc.styles['Caption'].paragraph_format.space_after=Pt(10)
doc.core_properties.author='Aruma Harada';doc.core_properties.title='Accuracy and representation tradeoffs in track record selected councils'
doc.core_properties.subject='PKD computational mechanism study';doc.core_properties.keywords='agent-based simulation, collective estimation, sortition, PKD'
footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=el('w:fldSimple',**{'w:instr':'PAGE'});footer._p.append(field)
tables=json.loads(Path('paper/tables.json').read_text())

def make_table(key,caption):
    p=doc.add_paragraph(caption,'Caption');p.paragraph_format.keep_with_next=True
    spec=tables[key];t=doc.add_table(rows=1,cols=len(spec['headers']));t.autofit=False
    t.style='Table Grid'
    borders=el('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        borders.append(el('w:'+side,**{'w:val':'single','w:sz':'4','w:color':'D9D9D9'}))
    t._tbl.tblPr.append(borders)
    for c,width in zip(t.columns,spec['widths']):c.width=Inches(width)
    t.rows[0]._tr.get_or_add_trPr().append(el('w:tblHeader'))
    for idx,rowvals in enumerate([spec['headers']]+spec['rows']):
        row=t.rows[0] if idx==0 else t.add_row()
        row._tr.get_or_add_trPr().append(el('w:cantSplit'))
        for j,(c,value,width) in enumerate(zip(row.cells,rowvals,spec['widths'])):
            c.width=Inches(width);p=c.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.02
            p.paragraph_format.keep_with_next=(idx==0)
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 or key in ['conditions','parameters'] else WD_ALIGN_PARAGRAPH.CENTER
            r=p.add_run(str(value));r.font.size=Pt(9.5 if key=='conditions' else 10)
            if idx==0:r.bold=True;c._tc.get_or_add_tcPr().append(el('w:shd',**{'w:fill':'E9E9E9'}))
            margins=el('w:tcMar')
            for side in ['top','bottom','left','right']:margins.append(el('w:'+side,**{'w:w':'70','w:type':'dxa'}))
            c._tc.get_or_add_tcPr().append(margins)
            c._tc.get_or_add_tcPr().append(el('w:vAlign',**{'w:val':'center'}))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

front=True
for line in Path('paper/manuscript_source.md').read_text(encoding='utf-8').splitlines():
    if not line:continue
    if line.startswith('# '):doc.add_paragraph(line[2:],'Title');continue
    if line.startswith('## '):
        front=False;p=doc.add_paragraph(line[3:],'Heading 1')
        if line.startswith('## References') or line.startswith('## Appendix A'):p.paragraph_format.page_break_before=True
        continue
    if line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2');continue
    if line.startswith('[[EQ|'):
        n=line[5:-2];p=doc.add_paragraph();p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(9)
        p.paragraph_format.keep_together=True;p.paragraph_format.tab_stops.add_tab_stop(Inches(6.7),WD_TAB_ALIGNMENT.RIGHT)
        math=el('m:oMath')
        for item in eq(n):math.append(item)
        p._p.append(math);p.add_run('\t('+n+')');continue
    if line.startswith('[[FIG|'):
        _,name,caption=line[2:-2].split('|',2)
        p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(8)
        shape=p.add_run().add_picture(str(Path('results/preprint_2026/analysis')/(name+'.png')),width=Inches(6.7));shape._inline.docPr.set('descr',caption)
        doc.add_paragraph(caption,'Caption');continue
    if line.startswith('[[TABLE|'):
        _,key,caption=line[2:-2].split('|',2);make_table(key,caption);continue
    if line=='[[REFERENCES]]':
        for ref in json.loads(Path('paper/references.json').read_text()):
            p=doc.add_paragraph();p.paragraph_format.left_indent=Inches(.2);p.paragraph_format.first_line_indent=Inches(-.2);p.paragraph_format.keep_together=True
            inline(p,ref)
        continue
    p=doc.add_paragraph();inline(p,line)
    if front:p.paragraph_format.space_after=Pt(3)
parser=argparse.ArgumentParser();parser.add_argument('--output',type=Path,default=Path('paper/PKD_preprint_v1.docx'));args=parser.parse_args()
dest=args.output;dest.parent.mkdir(exist_ok=True)
doc.save(dest)
print(dest.resolve())
